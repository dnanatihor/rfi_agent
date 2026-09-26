from __future__ import annotations

import csv
import io
import re
from dataclasses import dataclass
from pathlib import Path

from rfi_agent.ingest.files import parse_file

QUESTION_HEADERS = {"question", "questions", "query", "item", "prompt", "requirement", "rfi question"}
CATEGORY_HEADERS = {"category", "domain", "section", "topic"}
CONSTRAINT_HEADERS = {"constraints", "guidance", "format", "limit", "notes", "expected", "instructions"}
SKIP_HEADERS = {"#", "no", "number", "id", "n"}
NUMBERED = re.compile(r"^\s*\d+[\.\)]\s+(.+)$")


@dataclass
class ParsedQuestion:
    ordinal: int
    text: str
    category: str = ""
    constraints: str = ""


def parse_questionnaire(filename: str, data: bytes, *, max_questions: int = 500) -> list[ParsedQuestion]:
    name = filename.lower()
    if name.endswith(".csv"):
        rows = _from_tabular(_csv_rows(data))
    elif name.endswith(".xlsx"):
        rows = _from_tabular(_xlsx_rows(data))
    elif name.endswith(".docx"):
        rows = _from_docx(data)
    else:
        text = parse_file(filename, data).text
        rows = _from_lines(text.splitlines())
    cleaned = [q for q in rows if q.text.strip()]
    if len(cleaned) > max_questions:
        raise ValueError(f"Questionnaire exceeds the {max_questions} question cap.")
    if not cleaned:
        raise ValueError("No questions found in the file.")
    for i, item in enumerate(cleaned, start=1):
        item.ordinal = i
    return cleaned


def write_questionnaire_file(root: Path, session_id: str, filename: str, data: bytes) -> Path:
    dest = root / "questionnaires" / session_id
    dest.mkdir(parents=True, exist_ok=True)
    path = dest / Path(filename).name
    path.write_bytes(data)
    return path


def _csv_rows(data: bytes) -> list[list[str]]:
    text = data.decode("utf-8-sig", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return [[(c or "").strip() for c in row] for row in reader]


def _xlsx_rows(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    sheet = wb.worksheets[0]
    rows: list[list[str]] = []
    for row in sheet.iter_rows(values_only=True):
        cells = ["" if c is None else str(c).strip() for c in row]
        if any(cells):
            rows.append(cells)
    return rows


def _header_index(header: list[str], names: set[str]) -> int | None:
    lowered = [h.strip().lower() for h in header]
    for i, name in enumerate(lowered):
        if name in names:
            return i
    return None


def _from_tabular(rows: list[list[str]]) -> list[ParsedQuestion]:
    if not rows:
        return []
    header = rows[0]
    if _looks_like_header(header):
        q_idx = _header_index(header, QUESTION_HEADERS)
        if q_idx is None:
            q_idx = next((i for i, h in enumerate(header) if h.lower() not in SKIP_HEADERS and h != "#"), 0)
        cat_idx = _header_index(header, CATEGORY_HEADERS)
        con_idx = _header_index(header, CONSTRAINT_HEADERS)
        data_rows = rows[1:]
    else:
        q_idx = 1 if header and header[0].replace(".", "").isdigit() and len(header) > 1 else 0
        cat_idx = q_idx + 1 if len(header) > q_idx + 1 else None
        con_idx = q_idx + 2 if len(header) > q_idx + 2 else None
        data_rows = rows
    out: list[ParsedQuestion] = []
    for row in data_rows:
        if q_idx >= len(row):
            continue
        text = row[q_idx].strip()
        if not text or text.lower() in QUESTION_HEADERS:
            continue
        category = row[cat_idx].strip() if cat_idx is not None and cat_idx < len(row) else ""
        constraints = row[con_idx].strip() if con_idx is not None and con_idx < len(row) else ""
        out.append(ParsedQuestion(ordinal=len(out) + 1, text=text, category=category, constraints=constraints))
    return out


def _looks_like_header(header: list[str]) -> bool:
    joined = " ".join(h.lower() for h in header)
    return any(word in joined for word in QUESTION_HEADERS | CATEGORY_HEADERS | SKIP_HEADERS)


def _from_docx(data: bytes) -> list[ParsedQuestion]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    if doc.tables:
        rows: list[list[str]] = []
        for table in doc.tables:
            for row in table.rows:
                rows.append([(cell.text or "").strip() for cell in row.cells])
        parsed = _from_tabular(rows)
        if parsed:
            return parsed
    return _from_lines(p.text for p in doc.paragraphs)


def _from_lines(lines) -> list[ParsedQuestion]:
    out: list[ParsedQuestion] = []
    for line in lines:
        raw = (line or "").strip()
        if not raw:
            continue
        match = NUMBERED.match(raw)
        text = match.group(1).strip() if match else raw
        if len(text) < 8:
            continue
        out.append(ParsedQuestion(ordinal=len(out) + 1, text=text))
    return out
