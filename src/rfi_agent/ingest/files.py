from __future__ import annotations

from dataclasses import dataclass, field
from io import BytesIO
from pathlib import Path

from rfi_agent.ingest.chunker import chunk_text


@dataclass
class ParsedFile:
    text: str
    tables: list[list[list[str]]] = field(default_factory=list)
    content_type: str = "text/plain"
    page_count: int = 1
    image_only: bool = False


def parse_bytes(filename: str, data: bytes) -> str:
    return parse_file(filename, data).text


def parse_file(filename: str, data: bytes) -> ParsedFile:
    name = filename.lower()
    if name.endswith(".md"):
        return ParsedFile(text=data.decode("utf-8", errors="replace"), content_type="text/markdown")
    if name.endswith(".txt"):
        return ParsedFile(text=data.decode("utf-8", errors="replace"), content_type="text/plain")
    if name.endswith(".html") or name.endswith(".htm"):
        return ParsedFile(text=_strip_html(data.decode("utf-8", errors="replace")), content_type="text/html")
    if name.endswith(".pdf"):
        return _pdf(data)
    if name.endswith(".docx"):
        return _docx(data)
    if name.endswith(".pptx"):
        return _pptx(data)
    if name.endswith(".xlsx") or name.endswith(".csv"):
        return _xlsx(data, name)
    return ParsedFile(text=data.decode("utf-8", errors="replace"))


def file_chunks(filename: str, data: bytes) -> list[str]:
    return chunk_text(parse_bytes(filename, data))


def write_upload(root: Path, session_id: str, filename: str, data: bytes) -> Path:
    dest_dir = root / "uploads" / session_id
    dest_dir.mkdir(parents=True, exist_ok=True)
    path = dest_dir / Path(filename).name
    path.write_bytes(data)
    return path


def _pdf(data: bytes) -> ParsedFile:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    pages = []
    for i, page in enumerate(reader.pages, start=1):
        text = page.extract_text() or ""
        if text.strip():
            pages.append(f"# Page {i}\n{text}")
    if not pages:
        return ParsedFile(text="", content_type="application/pdf", page_count=len(reader.pages), image_only=True)
    return ParsedFile(
        text="\n\n".join(pages),
        content_type="application/pdf",
        page_count=len(reader.pages),
    )


def _docx(data: bytes) -> ParsedFile:
    from docx import Document

    doc = Document(BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text]
    tables: list[list[list[str]]] = []
    for table in doc.tables:
        grid = [[(cell.text or "").strip() for cell in row.cells] for row in table.rows]
        tables.append(grid)
        parts.append(_table_as_text(grid))
    return ParsedFile(
        text="\n".join(parts),
        tables=tables,
        content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        page_count=max(1, len(doc.paragraphs)),
    )


def _pptx(data: bytes) -> ParsedFile:
    from pptx import Presentation

    pres = Presentation(BytesIO(data))
    slides = []
    for i, slide in enumerate(pres.slides, start=1):
        bits = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False):
                bits.append(shape.text_frame.text)
        body = "\n".join(b for b in bits if b and b.strip())
        if body.strip():
            slides.append(f"# Slide {i}\n{body}")
    return ParsedFile(
        text="\n\n".join(slides),
        content_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
        page_count=len(pres.slides),
    )


def _xlsx(data: bytes, name: str) -> ParsedFile:
    if name.endswith(".csv"):
        text = data.decode("utf-8", errors="replace")
        rows = [line.split(",") for line in text.splitlines() if line.strip()]
        return ParsedFile(text=text, tables=[rows] if rows else [], content_type="text/csv", page_count=1)
    from openpyxl import load_workbook

    wb = load_workbook(BytesIO(data), read_only=True, data_only=True)
    parts = []
    tables: list[list[list[str]]] = []
    for sheet in wb.worksheets:
        parts.append(f"# {sheet.title}")
        grid: list[list[str]] = []
        for row in sheet.iter_rows(values_only=True):
            cells = ["" if c is None else str(c) for c in row]
            if any(cells):
                grid.append(cells)
                parts.append(" | ".join(cells))
        if grid:
            tables.append(grid)
    return ParsedFile(
        text="\n".join(parts),
        tables=tables,
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        page_count=len(wb.worksheets),
    )


def _table_as_text(grid: list[list[str]]) -> str:
    if not grid:
        return ""
    lines = ["# Table"]
    for row in grid:
        lines.append(" | ".join(row))
    return "\n".join(lines)


def _strip_html(html: str) -> str:
    import re

    text = re.sub(r"(?is)<script.*?>.*?</script>", " ", html)
    text = re.sub(r"(?is)<style.*?>.*?</style>", " ", text)
    text = re.sub(r"(?s)<[^>]+>", " ", text)
    return re.sub(r"\s+", " ", text).strip()
